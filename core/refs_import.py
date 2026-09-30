"""Загрузка справочников заказчика из одного Excel-файла (шаблон: manage.py import_refs --template).

Листы (порядок важен — ссылки идут на предыдущие листы):
  Подразделения, Пользователи, Статьи бюджета, Категории, Номенклатура, Поставщики,
  Рамочные договоры, Пороги согласования.

Повторная загрузка безопасна: записи обновляются по ключу (код / логин / БИН / наименование),
новые — добавляются, ничего не удаляется. Режим проверки (dry_run) ничего не записывает.
"""
import io
import re
import secrets
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.contrib.auth.models import Group, User
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from . import roles
from .models import (
    ApprovalRule, BudgetItem, Category, Contract, Department, Nomenclature, Profile, Supplier,
)

# Лист → [(колонка, обязательна, пояснение, пример)]
SHEETS = {
    "Подразделения": [
        ("Наименование", True, "Уникальное название", "IT-отдел"),
        ("Руководитель (логин)", False, "Согласует СЗ подразделения. Логин с листа «Пользователи»", "saparova"),
    ],
    "Пользователи": [
        ("Логин", True, "Латиница, без пробелов. Лучше как в AD/почте", "ivanov"),
        ("Фамилия", True, "", "Иванов"),
        ("Имя", True, "", "Иван"),
        ("E-mail", False, "Для уведомлений", "ivanov@company.kz"),
        ("Подразделение", False, "Точно как на листе «Подразделения»", "IT-отдел"),
        ("Должность", False, "", "Системный администратор"),
        ("Роли", True, "Через запятую: " + ", ".join(roles.ALL_ROLES), "Инициатор"),
    ],
    "Статьи бюджета": [
        ("Код", True, "Как в бюджете / 1С", "01.01"),
        ("Наименование", True, "", "Канцелярские расходы"),
    ],
    "Категории": [
        ("Наименование", True, "Группа закупок; по ней подбираются поставщики", "Канцтовары"),
    ],
    "Номенклатура": [
        ("Наименование", True, "", "Бумага А4 SvetoCopy, 500 л."),
        ("Ед. изм.", True, "шт, пач, уп, кг, л, м, усл …", "пач"),
        ("Категория", False, "Как на листе «Категории»", "Канцтовары"),
        ("Код 1С", False, "Код номенклатуры в 1С — ключ сопоставления поступлений", "00-0001"),
    ],
    "Поставщики": [
        ("Наименование", True, "", "ТОО «Канцлер»"),
        ("БИН/ИИН", False, "12 цифр. Ключ поиска при повторной загрузке", "180540012345"),
        ("E-mail", False, "Куда отправлять запросы КП", "sales@kancler.kz"),
        ("Телефон", False, "", "+7 727 000 00 00"),
        ("Контактное лицо", False, "", "Асель"),
        ("Категории", False, "Через точку с запятой", "Канцтовары; Хозтовары"),
        ("Код 1С", False, "Код контрагента в 1С", "000123"),
    ],
    "Рамочные договоры": [
        ("Номер", True, "", "РД-2026-001"),
        ("Дата", True, "ДД.ММ.ГГГГ", "15.01.2026"),
        ("Поставщик (БИН или наименование)", True, "", "180540012345"),
        ("Действует до", False, "ДД.ММ.ГГГГ", "31.12.2026"),
    ],
    "Пороги согласования": [
        ("Сумма от", True, "Сумма решения по закупке, ₸", "1000000"),
        ("Роль", True, "Директор / Финдиректор", "Директор"),
    ],
}


class ImportReport:
    def __init__(self):
        self.created = {}
        self.updated = {}
        self.errors = []
        self.passwords = []  # (логин, временный пароль) — только для новых пользователей

    def add(self, sheet, created):
        d = self.created if created else self.updated
        d[sheet] = d.get(sheet, 0) + 1

    def error(self, sheet, row, text):
        self.errors.append(f"{sheet}, строка {row}: {text}")

    @property
    def ok(self):
        return not self.errors


# ---------------------------------------------------------------- шаблон


def build_template():
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    req_fill = PatternFill("solid", fgColor="B42318")
    info = wb.create_sheet("Инструкция")
    lines = [
        ("Как заполнить", True),
        ("1. Заполните листы по порядку, начиная со второй строки. Красные колонки (*) обязательны.", False),
        ("2. Подсказка к колонке — в примечании к заголовку (наведите мышь).", False),
        ("3. Названия подразделений, категорий и логины должны совпадать между листами символ в символ.", False),
        ("4. Не меняйте названия листов и колонок. Пустой лист просто не загружается.", False),
        ("5. Файл можно загружать повторно: существующие записи обновятся, новые добавятся, ничего не удалится.", False),
        ("", False),
        ("Роли: " + ", ".join(roles.ALL_ROLES), False),
        ("", False),
        ("Примеры заполнения", True),
    ]
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        for i, (col, required, hint, _) in enumerate(cols, 1):
            c = ws.cell(row=1, column=i, value=col + (" *" if required else ""))
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = req_fill if required else head_fill
            if hint:
                c.comment = Comment(hint, "Шаблон")
            ws.column_dimensions[c.column_letter].width = max(16, len(col) + 6)
        ws.freeze_panes = "A2"
        lines.append((f"{name}: " + " | ".join(f"{c}: {ex}" for c, _, _, ex in cols), False))
    for i, (t, bold) in enumerate(lines, 1):
        info.cell(row=i, column=1, value=t).font = Font(bold=bold, size=13 if bold else 11)
    info.column_dimensions["A"].width = 140
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- чтение


def _norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _rows(wb, sheet):
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    try:
        header = [_norm(h).rstrip(" *") for h in next(it)]
    except StopIteration:
        return []
    out = []
    for n, row in enumerate(it, start=2):
        if not any(_norm(v) for v in row):
            continue
        rec = {header[i]: row[i] for i in range(min(len(header), len(row))) if header[i]}
        out.append((n, rec))
    return out


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _norm(v)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"не распознана дата «{s}», нужно ДД.ММ.ГГГГ")


def _decimal(v):
    s = _norm(v).replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        raise ValueError(f"не число «{v}»")


# ---------------------------------------------------------------- загрузка


def import_workbook(content, dry_run=False):
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    report = ImportReport()
    try:
        with transaction.atomic():
            _import(wb, report)
            if dry_run or report.errors:
                raise _Rollback()
    except _Rollback:
        if not report.errors:
            report.passwords = [(u, "(будет создан при загрузке)") for u, _ in report.passwords]
    return report


class _Rollback(Exception):
    pass


def _required(rec, cols, report, sheet, n):
    missing = [c for c, req, *_ in cols if req and not _norm(rec.get(c))]
    if missing:
        report.error(sheet, n, "не заполнено: " + ", ".join(missing))
        return False
    return True


def _import(wb, report):
    for group in roles.ALL_ROLES:
        Group.objects.get_or_create(name=group)
    heads = []

    sheet = "Подразделения"
    for n, r in _rows(wb, sheet):
        if not _required(r, SHEETS[sheet], report, sheet, n):
            continue
        dept, created = Department.objects.get_or_create(name=_norm(r["Наименование"]))
        report.add(sheet, created)
        if _norm(r.get("Руководитель (логин)")):
            heads.append((n, dept, _norm(r["Руководитель (логин)"]).lower()))

    sheet = "Пользователи"
    for n, r in _rows(wb, sheet):
        if not _required(r, SHEETS[sheet], report, sheet, n):
            continue
        login = _norm(r["Логин"]).lower()
        if not re.fullmatch(r"[a-z0-9_.@+-]+", login):
            report.error(sheet, n, f"логин «{login}»: только латиница, цифры и . _ - @")
            continue
        role_names = [x.strip() for x in re.split(r"[,;]", _norm(r["Роли"])) if x.strip()]
        bad = [x for x in role_names if x not in roles.ALL_ROLES]
        if bad:
            report.error(sheet, n, f"неизвестные роли: {', '.join(bad)}")
            continue
        dept = None
        if _norm(r.get("Подразделение")):
            dept = Department.objects.filter(name=_norm(r["Подразделение"])).first()
            if not dept:
                report.error(sheet, n, f"подразделение «{_norm(r['Подразделение'])}» не найдено")
                continue
        user, created = User.objects.get_or_create(username=login)
        user.last_name, user.first_name = _norm(r["Фамилия"]), _norm(r["Имя"])
        user.email = _norm(r.get("E-mail"))
        user.is_staff = roles.ADMIN in role_names
        if created:
            password = secrets.token_urlsafe(9)
            user.set_password(password)
            report.passwords.append((login, password))
        user.save()
        user.groups.set(Group.objects.filter(name__in=role_names))
        Profile.objects.update_or_create(user=user, defaults={"department": dept, "position": _norm(r.get("Должность"))})
        report.add(sheet, created)

    for n, dept, login in heads:
        head = User.objects.filter(username=login).first()
        if not head:
            report.error("Подразделения", n, f"руководитель «{login}» не найден на листе «Пользователи»")
            continue
        dept.head = head
        dept.save(update_fields=["head"])

    sheet = "Статьи бюджета"
    for n, r in _rows(wb, sheet):
        if not _required(r, SHEETS[sheet], report, sheet, n):
            continue
        _, created = BudgetItem.objects.update_or_create(code=_norm(r["Код"]), defaults={"name": _norm(r["Наименование"])})
        report.add(sheet, created)

    sheet = "Категории"
    for n, r in _rows(wb, sheet):
        if not _required(r, SHEETS[sheet], report, sheet, n):
            continue
        _, created = Category.objects.get_or_create(name=_norm(r["Наименование"]))
        report.add(sheet, created)

    def category(name, sheet, n):
        if not name:
            return None
        c = Category.objects.filter(name=name).first()
        if not c:
            report.error(sheet, n, f"категория «{name}» не найдена на листе «Категории»")
        return c

    sheet = "Номенклатура"
    for n, r in _rows(wb, sheet):
        if not _required(r, SHEETS[sheet], report, sheet, n):
            continue
        name, code = _norm(r["Наименование"]), _norm(r.get("Код 1С"))
        cat = category(_norm(r.get("Категория")), sheet, n)
        obj = Nomenclature.objects.filter(code_1c=code).first() if code else None
        obj = obj or Nomenclature.objects.filter(name=name).first()
        created = obj is None
        obj = obj or Nomenclature(name=name)
        obj.name, obj.unit, obj.category, obj.code_1c = name, _norm(r["Ед. изм."]), cat, code
        obj.save()
        report.add(sheet, created)

    sheet = "Поставщики"
    for n, r in _rows(wb, sheet):
        if not _required(r, SHEETS[sheet], report, sheet, n):
            continue
        name, bin_ = _norm(r["Наименование"]), re.sub(r"\D", "", _norm(r.get("БИН/ИИН")))
        if bin_ and len(bin_) != 12:
            report.error(sheet, n, f"БИН/ИИН «{bin_}» должен состоять из 12 цифр")
            continue
        obj = Supplier.objects.filter(bin=bin_).first() if bin_ else None
        obj = obj or Supplier.objects.filter(name=name).first()
        created = obj is None
        obj = obj or Supplier(name=name)
        obj.name, obj.bin = name, bin_
        obj.email, obj.phone = _norm(r.get("E-mail")), _norm(r.get("Телефон"))
        obj.contact, obj.code_1c = _norm(r.get("Контактное лицо")), _norm(r.get("Код 1С"))
        obj.save()
        cats = [category(x.strip(), sheet, n) for x in _norm(r.get("Категории")).split(";") if x.strip()]
        obj.categories.set([c for c in cats if c])
        report.add(sheet, created)

    sheet = "Рамочные договоры"
    for n, r in _rows(wb, sheet):
        if not _required(r, SHEETS[sheet], report, sheet, n):
            continue
        key = _norm(r["Поставщик (БИН или наименование)"])
        supplier = Supplier.objects.filter(bin=re.sub(r"\D", "", key)).first() if re.sub(r"\D", "", key) else None
        supplier = supplier or Supplier.objects.filter(name=key).first()
        if not supplier:
            report.error(sheet, n, f"поставщик «{key}» не найден")
            continue
        try:
            d = _date(r["Дата"])
            until = _date(r["Действует до"]) if _norm(r.get("Действует до")) else None
        except ValueError as e:
            report.error(sheet, n, str(e))
            continue
        _, created = Contract.objects.update_or_create(
            number=_norm(r["Номер"]), supplier=supplier, kind=Contract.Kind.FRAMEWORK,
            defaults={"date": d, "valid_until": until, "status": Contract.Status.SIGNED},
        )
        report.add(sheet, created)

    sheet = "Пороги согласования"
    for n, r in _rows(wb, sheet):
        if not _required(r, SHEETS[sheet], report, sheet, n):
            continue
        role = _norm(r["Роль"])
        if role not in roles.ALL_ROLES:
            report.error(sheet, n, f"неизвестная роль «{role}»")
            continue
        try:
            amount = _decimal(r["Сумма от"])
        except ValueError as e:
            report.error(sheet, n, str(e))
            continue
        _, created = ApprovalRule.objects.get_or_create(min_amount=amount, role=role)
        report.add(sheet, created)
