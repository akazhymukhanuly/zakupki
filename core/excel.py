"""Excel: выгрузка таблиц (сравнение КП, отчёты, 1С) и шаблон для импорта КП."""
import io
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BEST_FILL = PatternFill("solid", fgColor="C6EFCE")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _cell_value(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


def write_sheet(ws, rows, title=None, widths=None):
    start = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        start = 3
    for r_i, row in enumerate(rows):
        for c_i, v in enumerate(row, 1):
            cell = ws.cell(row=start + r_i, column=c_i, value=_cell_value(v))
            cell.border = BORDER
            if r_i == 0:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            elif isinstance(v, (date, datetime)):
                cell.number_format = "DD.MM.YYYY"
            elif isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
                cell.number_format = "#,##0.00" if isinstance(v, (float, Decimal)) else "0"
    for c_i in range(1, (max(len(r) for r in rows) if rows else 0) + 1):
        width = (widths or {}).get(c_i)
        if width is None:
            width = min(max(len(str(r[c_i - 1])) if c_i - 1 < len(r) and r[c_i - 1] is not None else 0 for r in rows) + 2, 60)
        ws.column_dimensions[get_column_letter(c_i)].width = max(width, 8)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    return start


def response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{_quote(filename)}"
    return resp


def _quote(s):
    from urllib.parse import quote
    return quote(s)


def rows_response(rows, filename, title=None, sheet="Лист1"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    write_sheet(ws, rows, title)
    return response(wb, filename)


def comparison_workbook(proc, cmp):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение КП"
    ws.cell(row=1, column=1, value=f"Сравнительная таблица КП — закупка №{proc.number} «{proc.title}»").font = Font(bold=True, size=13)
    header = ["Позиция", "СЗ", "Наименование", "Кол-во", "Ед."]
    for t in cmp["totals"]:
        header += [f"{t['quote'].supplier.name}\nцена", "срок, дн.", "сумма"]
    header += ["Мин. цена"]
    rows = [header]
    for row in cmp["rows"]:
        line = row["line"]
        r = [line.item.code, f"№{line.item.memo.number}", line.item.description, float(line.quantity), line.item.unit]
        for c in row["cells"]:
            ql = c["ql"]
            if not ql:
                r += ["—", None, None]
            elif ql.not_offered:
                r += ["не предлагает", None, None]
            else:
                price = float(ql.price)
                label = price if not ql.is_analog else f"{price:,.2f} (аналог: {ql.analog_description})"
                r += [label, ql.lead_time_days, float(ql.price * line.quantity)]
        r.append(float(row["min_price"]) if row["min_price"] is not None else None)
        rows.append(r)
    total = ["ИТОГО", "", "", None, ""]
    for t in cmp["totals"]:
        total += ["", "", float(t["total"])]
    total.append(None)
    rows.append(total)
    rows.append(["Итого «лучшее по каждой строке»", "", "", None, ""] + [""] * (3 * len(cmp["totals"])) + [float(cmp["best_total"])])
    start = write_sheet(ws, rows, widths={3: 40})
    # Подсветка минимальной цены по строке.
    for r_i, row in enumerate(cmp["rows"], start=start + 1):
        for q_i, c in enumerate(row["cells"]):
            if c["is_min"]:
                ws.cell(row=r_i, column=6 + q_i * 3).fill = BEST_FILL
    for c_i in range(1, len(header) + 1):
        ws.cell(row=start + len(rows) - 2, column=c_i).font = Font(bold=True)
        ws.cell(row=start + len(rows) - 1, column=c_i).font = Font(bold=True)
    return wb


# ---------------------------------------------------------------- шаблон КП


QUOTE_TEMPLATE_HEADER = ["ID строки", "Позиция", "Наименование", "Кол-во", "Ед.", "Цена за ед.",
                         "Срок поставки, дн.", "Не предлагает (да/нет)", "Аналог: описание", "Предлагаемое кол-во", "Комментарий"]


def quote_template(proc, rfq, lines):
    wb = Workbook()
    ws = wb.active
    ws.title = "КП"
    rows = [QUOTE_TEMPLATE_HEADER]
    for l in lines:
        rows.append([l.pk, l.item.code, l.item.description, float(l.quantity), l.item.unit, None, None, None, None, None, None])
    write_sheet(ws, rows, title=f"КП {rfq.supplier.name} — закупка №{proc.number}", widths={3: 45, 9: 30, 11: 30})
    return wb


def parse_quote_workbook(uploaded):
    wb = load_workbook(io.BytesIO(uploaded.read()), data_only=True, read_only=True)
    ws = wb.active
    result = {}
    header_found = False
    for row in ws.iter_rows(values_only=True):
        if not header_found:
            if row and row[0] == QUOTE_TEMPLATE_HEADER[0]:
                header_found = True
            continue
        if not row or row[0] in (None, ""):
            continue
        try:
            line_id = int(row[0])
        except (TypeError, ValueError):
            continue

        def dec(v):
            if v in (None, ""):
                return None
            return Decimal(str(v).replace(",", ".").replace(" ", "").replace("\xa0", ""))

        not_offered = str(row[7] or "").strip().lower() in ("да", "yes", "1", "true", "+")
        analog = str(row[8] or "").strip()
        result[line_id] = {
            "price": dec(row[5]),
            "lead_time_days": int(row[6]) if row[6] not in (None, "") else None,
            "not_offered": not_offered,
            "is_analog": bool(analog),
            "analog_description": analog,
            "offered_qty": dec(row[9]),
            "comment": str(row[10] or ""),
        }
    if not header_found:
        from .services import BusinessError
        raise BusinessError("Файл не похож на шаблон КП: скачайте шаблон на этой странице.")
    return result
