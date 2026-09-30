"""Обмен с 1С (п. 8 ТЗ) через файлы.

Выгрузка: договор со спецификацией построчно, в каждой строке — ID позиции («СЗ-12/3»),
чтобы поступление можно было разнести обратно. Форматы: XML, JSON, Excel.

Загрузка: поступления ТМЦ/акты (построчно) и оплаты (по договору целиком)
из Excel/CSV/XML/JSON. Сопоставление строк поступления:
  1) по ID позиции — подтверждается автоматически;
  2) по номенклатуре (код 1С или наименование) в рамках договора — требует ручного подтверждения;
  3) не найдено — строка остаётся несопоставленной, закупщик разносит вручную.

Файловый обмен выбран для MVP как самый простой для 1С-программиста: обработка
в 1С читает/пишет эти файлы стандартными средствами (ЧтениеXML / ЧтениеJSON / табличный документ).
Позже тот же формат можно отдавать через HTTP-сервис (REST) без изменения логики.
"""
import csv
import io
import json
import re
import xml.etree.ElementTree as ET
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.utils import timezone

from .models import Contract, ContractLine, MemoItem, ReceiptLine
from .services import BusinessError, register_payment, register_receipt

FORMAT_VERSION = "1.0"

RECEIPT_COLUMNS = ["Договор", "Дата", "НомерДокумента", "IDПозиции", "Номенклатура", "КодНоменклатуры", "Количество"]
PAYMENT_COLUMNS = ["Договор", "Дата", "НомерДокумента", "Сумма"]

ALIASES = {
    "договор": "Договор", "номердоговора": "Договор", "contract": "Договор",
    "дата": "Дата", "date": "Дата",
    "номердокумента": "НомерДокумента", "документ": "НомерДокумента", "doc": "НомерДокумента", "docnumber": "НомерДокумента",
    "idпозиции": "IDПозиции", "idпозиция": "IDПозиции", "позиция": "IDПозиции", "positionid": "IDПозиции", "id": "IDПозиции",
    "номенклатура": "Номенклатура", "наименование": "Номенклатура", "name": "Номенклатура",
    "кодноменклатуры": "КодНоменклатуры", "код": "КодНоменклатуры", "code": "КодНоменклатуры",
    "количество": "Количество", "колво": "Количество", "qty": "Количество", "quantity": "Количество",
    "сумма": "Сумма", "amount": "Сумма",
}


# ---------------------------------------------------------------- выгрузка


def contract_payload(contract, collapse=None):
    if collapse is None:
        collapse = settings.PROCUREMENT["ONEC_COLLAPSE_CONSOLIDATED"]
    lines = list(contract.lines.select_related("item", "item__memo", "item__nomenclature"))
    rows = []
    if collapse:
        grouped = OrderedDict()
        for l in lines:
            key = (l.item.nomenclature_id or f"l{l.pk}", l.price)
            grouped.setdefault(key, []).append(l)
        for group in grouped.values():
            first = group[0]
            rows.append({
                "ids": [l.item.code for l in group],
                "line_ids": [l.pk for l in group],
                "name": first.item.nomenclature.name if first.item.nomenclature else first.description,
                "code_1c": first.item.nomenclature.code_1c if first.item.nomenclature else "",
                "unit": first.unit,
                "quantity": sum((l.quantity for l in group), Decimal(0)),
                "price": first.price,
                "memos": sorted({f"СЗ-{l.item.memo.number}" for l in group}),
            })
    else:
        for l in lines:
            rows.append({
                "ids": [l.item.code], "line_ids": [l.pk], "name": l.description,
                "code_1c": l.item.nomenclature.code_1c if l.item.nomenclature else "",
                "unit": l.unit, "quantity": l.quantity, "price": l.price, "memos": [f"СЗ-{l.item.memo.number}"],
            })
    for r in rows:
        r["amount"] = r["quantity"] * r["price"]
        r["position_id"] = ";".join(r["ids"])
        r["comment"] = ", ".join(r["ids"])
    return {
        "number": contract.number,
        "date": contract.date,
        "kind": contract.get_kind_display(),
        "parent": contract.parent.number if contract.parent else "",
        "supplier": contract.supplier,
        "amount": sum((r["amount"] for r in rows), Decimal(0)),
        "rows": rows,
    }


def export_xml(contracts, collapse=None):
    root = ET.Element("ОбменСКП", ВерсияФормата=FORMAT_VERSION, ДатаВыгрузки=timezone.now().isoformat(timespec="seconds"))
    for c in contracts:
        p = contract_payload(c, collapse)
        el = ET.SubElement(root, "Договор", Номер=p["number"], Дата=p["date"].isoformat(), Вид=p["kind"],
                           Сумма=f"{p['amount']:.2f}", Валюта="KZT")
        if p["parent"]:
            el.set("РамочныйДоговор", p["parent"])
        s = p["supplier"]
        ET.SubElement(el, "Контрагент", Наименование=s.name, БИН=s.bin or "", Код1С=s.code_1c or "")
        spec = ET.SubElement(el, "Спецификация")
        for i, r in enumerate(p["rows"], 1):
            ET.SubElement(spec, "Строка", НомерСтроки=str(i), IDПозиции=r["position_id"],
                          Номенклатура=r["name"], КодНоменклатуры=r["code_1c"], ЕдиницаИзмерения=r["unit"],
                          Количество=_num(r["quantity"]), Цена=f"{r['price']:.2f}", Сумма=f"{r['amount']:.2f}",
                          Комментарий=r["comment"])
    ET.indent(root) if hasattr(ET, "indent") else None
    return b'<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(root, encoding="utf-8")


def export_json(contracts, collapse=None):
    data = {"version": FORMAT_VERSION, "exported_at": timezone.now().isoformat(timespec="seconds"), "contracts": []}
    for c in contracts:
        p = contract_payload(c, collapse)
        data["contracts"].append({
            "number": p["number"], "date": p["date"].isoformat(), "kind": p["kind"], "framework": p["parent"],
            "currency": "KZT", "amount": float(p["amount"]),
            "supplier": {"name": p["supplier"].name, "bin": p["supplier"].bin, "code_1c": p["supplier"].code_1c},
            "lines": [{
                "line": i, "position_id": r["position_id"], "name": r["name"], "code_1c": r["code_1c"],
                "unit": r["unit"], "quantity": float(r["quantity"]), "price": float(r["price"]),
                "amount": float(r["amount"]), "comment": r["comment"],
            } for i, r in enumerate(p["rows"], 1)],
        })
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


def export_rows(contracts, collapse=None):
    header = ["Договор", "Дата", "Вид", "Рамочный договор", "Контрагент", "БИН", "№", "IDПозиции",
              "Номенклатура", "КодНоменклатуры", "Ед.", "Количество", "Цена", "Сумма", "Комментарий"]
    out = [header]
    for c in contracts:
        p = contract_payload(c, collapse)
        for i, r in enumerate(p["rows"], 1):
            out.append([p["number"], p["date"], p["kind"], p["parent"], p["supplier"].name, p["supplier"].bin, i,
                        r["position_id"], r["name"], r["code_1c"], r["unit"], float(r["quantity"]), float(r["price"]),
                        float(r["amount"]), r["comment"]])
    return out


def mark_exported(contracts):
    now = timezone.now()
    for c in contracts:
        c.exported_1c_at = now
        c.save(update_fields=["exported_1c_at"])


def _num(v):
    s = f"{Decimal(v):f}"
    return s.rstrip("0").rstrip(".") if "." in s else s


# ---------------------------------------------------------------- чтение файлов


def _norm_key(k):
    k = re.sub(r"[\s._\-№/]", "", str(k or "")).lower()
    return ALIASES.get(k, str(k).strip())


def read_rows(uploaded, kind):
    """Читает файл любого поддерживаемого формата в список словарей с нормализованными ключами."""
    name = uploaded.name.lower()
    content = uploaded.read()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [_norm_key(h) for h in next(it)]
        rows = [dict(zip(header, r)) for r in it if any(v not in (None, "") for v in r)]
    elif name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=";,\t")
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        rows = [{_norm_key(k): v for k, v in r.items()} for r in reader]
    elif name.endswith(".json"):
        data = json.loads(content.decode("utf-8-sig"))
        if isinstance(data, dict):
            data = data.get("rows") or data.get("receipts") or data.get("payments") or data.get("lines") or []
        rows = [{_norm_key(k): v for k, v in r.items()} for r in data]
    elif name.endswith(".xml"):
        root = ET.fromstring(content)
        rows = []
        for doc in root:
            base = {_norm_key(k): v for k, v in doc.attrib.items()}
            children = list(doc)
            if children:
                for ch in children:
                    r = dict(base)
                    r.update({_norm_key(k): v for k, v in ch.attrib.items()})
                    rows.append(r)
            else:
                rows.append(base)
    else:
        raise BusinessError("Поддерживаются файлы .xlsx, .csv, .json, .xml")
    required = RECEIPT_COLUMNS[:1] + ["Количество"] if kind == "receipt" else ["Договор", "Сумма"]
    if rows:
        missing = [c for c in required if c not in rows[0]]
        if missing:
            raise BusinessError(f"В файле нет колонок: {', '.join(missing)}. Ожидаются: "
                                f"{', '.join(RECEIPT_COLUMNS if kind == 'receipt' else PAYMENT_COLUMNS)}")
    return rows


def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%Y-%m-%dT%H:%M:%S", "%d.%m.%Y %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return timezone.localdate()


def _to_dec(v):
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = str(v or "").replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        raise BusinessError(f"Не число: «{v}»")


def _find_contract(number):
    number = str(number or "").strip()
    c = Contract.objects.filter(number=number).first()
    if not c:
        raise BusinessError(f"Договор «{number}» не найден.")
    return c


def _line_by_code(contract, code):
    code = code.strip()
    m = re.match(r"^СЗ-(\d+)/(\d+)$", code, re.I) or re.match(r"^SZ-(\d+)/(\d+)$", code, re.I)
    qs = contract.lines.select_related("item", "item__memo", "item__nomenclature")
    if m:
        return qs.filter(item__memo__number=int(m.group(1)), item__line_no=int(m.group(2))).first()
    if code.isdigit():
        return qs.filter(item_id=int(code)).first()
    return None


def _line_by_nomenclature(contract, name, code_1c):
    qs = contract.lines.select_related("item", "item__nomenclature")
    if code_1c:
        found = qs.filter(item__nomenclature__code_1c=str(code_1c).strip()).first()
        if found:
            return found
    name = str(name or "").strip().lower()
    if not name:
        return None
    for l in qs:
        candidates = [l.description.lower(), l.item.description.lower()]
        if l.item.nomenclature:
            candidates.append(l.item.nomenclature.name.lower())
        if any(name == c or name in c or c in name for c in candidates):
            return l
    return None


def import_receipts(uploaded, user):
    """Возвращает отчёт: сколько строк сопоставлено по ID, по номенклатуре (ждут подтверждения), не найдено."""
    rows = read_rows(uploaded, "receipt")
    docs = OrderedDict()
    for r in rows:
        key = (str(r.get("Договор", "")).strip(), str(r.get("НомерДокумента") or "").strip(), _to_date(r.get("Дата")))
        docs.setdefault(key, []).append(r)
    report = {"docs": 0, "by_id": 0, "by_nomenclature": 0, "unmatched": 0, "errors": []}
    for (number, doc_no, d), doc_rows in docs.items():
        try:
            contract = _find_contract(number)
            entries = []
            for r in doc_rows:
                qty = _to_dec(r.get("Количество"))
                raw_ref = str(r.get("IDПозиции") or "").strip()
                raw_name = str(r.get("Номенклатура") or "").strip()
                refs = [x for x in re.split(r"[;,]\s*", raw_ref) if x] if raw_ref else []
                lines = [l for l in (_line_by_code(contract, x) for x in refs) if l]
                if lines and len(lines) == len(refs):
                    # Свёрнутая строка (п. 6.1) — распределяем количество по строкам по порядку.
                    left = qty
                    for i, l in enumerate(lines):
                        take = left if i == len(lines) - 1 else min(left, max(l.quantity - l.delivered_qty, Decimal(0)))
                        if take > 0:
                            entries.append((l, take, ReceiptLine.Match.ID, True, raw_name, l.item.code))
                            report["by_id"] += 1
                        left -= take
                    continue
                l = _line_by_nomenclature(contract, raw_name, r.get("КодНоменклатуры"))
                if l:
                    entries.append((l, qty, ReceiptLine.Match.NOMENCLATURE, False, raw_name, raw_ref))
                    report["by_nomenclature"] += 1
                else:
                    entries.append((None, qty, ReceiptLine.Match.MANUAL, False, raw_name, raw_ref))
                    report["unmatched"] += 1
            register_receipt(contract, d, doc_no, entries, user)
            report["docs"] += 1
        except BusinessError as e:
            report["errors"].append(f"{number} / {doc_no}: {e}")
    return report


def import_payments(uploaded, user):
    rows = read_rows(uploaded, "payment")
    report = {"payments": 0, "errors": []}
    for r in rows:
        try:
            contract = _find_contract(r.get("Договор"))
            register_payment(contract, _to_date(r.get("Дата")), _to_dec(r.get("Сумма")),
                             str(r.get("НомерДокумента") or "").strip(), user)
            report["payments"] += 1
        except BusinessError as e:
            report["errors"].append(f"{r.get('Договор')}: {e}")
    return report


def sample_receipt_rows(contract=None):
    rows = [RECEIPT_COLUMNS]
    if contract:
        for l in contract.lines.select_related("item", "item__memo", "item__nomenclature"):
            rows.append([contract.number, timezone.localdate().strftime("%d.%m.%Y"), "ПТУ-0001", l.item.code,
                         l.description, l.item.nomenclature.code_1c if l.item.nomenclature else "", float(l.quantity)])
    return rows
