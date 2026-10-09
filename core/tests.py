"""Тесты правил ТЗ. Запуск: python manage.py test core"""
import io
import re
import json
from datetime import timedelta
from decimal import Decimal as D

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from . import onec, roles, services
from .models import (
    RFQ, Contract, Corridor, DecisionApproval, Station, DecisionLine, Department, Memo, MemoItem, Nomenclature, Procurement, ProcurementLine,
    ReceiptLine, Supplier, WithdrawalRequest,
)

S = MemoItem.Status


class Base(TestCase):
    def setUp(self):
        for r in roles.ALL_ROLES:
            Group.objects.get_or_create(name=r)
        self.initiator = self.mk("ini", roles.INITIATOR)
        self.initiator2 = self.mk("ini2", roles.INITIATOR)
        self.head = self.mk("head", roles.APPROVER)
        self.buyer = self.mk("buyer", roles.BUYER)
        self.director = self.mk("dir", roles.DIRECTOR)
        self.chief = self.mk("chief", roles.CHIEF)
        self.cfo = self.mk("cfo", roles.CFO)
        self.lawyer = self.mk("lawyer", roles.LAWYER)
        self.pto = self.mk("pto", roles.PTO)
        self.by_role = {roles.CHIEF: self.chief, roles.CFO: self.cfo, roles.BUYER: self.buyer,
                        roles.LAWYER: self.lawyer, roles.DIRECTOR: self.director, roles.PTO: self.pto}
        self.dept = Department.objects.create(name="Отдел", head=self.head)
        self.paper = Nomenclature.objects.create(name="Бумага А4", unit="пач", code_1c="P1")
        self.pens = Nomenclature.objects.create(name="Ручки", unit="уп", code_1c="P2")
        self.s1 = Supplier.objects.create(name="Поставщик 1")
        self.s2 = Supplier.objects.create(name="Поставщик 2")

    def mk(self, name, role):
        u = User.objects.create_user(name, password="x", first_name=name)
        u.groups.add(Group.objects.get(name=role))
        return u

    def memo(self, user, items, approve=True, rejected=None):
        m = Memo.objects.create(department=self.dept, initiator=user, justification="тест",
                                required_date=timezone.localdate() + timedelta(days=10))
        for n, (nom, qty) in enumerate(items, 1):
            MemoItem.objects.create(memo=m, line_no=n, nomenclature=nom, description=nom.name, quantity=D(qty),
                                    unit=nom.unit, required_date=m.required_date)
        services.submit_memo(m, user)
        if approve:
            services.approve_memo(m, self.head, rejected or {})
        return m

    def quote(self, proc, supplier, prices, offered=None, analog=None):
        rfq, _ = RFQ.objects.get_or_create(procurement=proc, supplier=supplier)
        rows = {}
        for line in proc.lines.filter(state="active"):
            nom = line.item.nomenclature
            if nom in prices:
                p = prices[nom]
                rows[line.pk] = {"price": None if p is None else D(p), "not_offered": p is None,
                                 "offered_qty": (offered or {}).get(nom),
                                 "is_analog": bool((analog or {}).get(nom)), "analog_description": (analog or {}).get(nom, "")}
        return services.save_quote(rfq, {}, rows, self.buyer)

    def approve_all(self, proc):
        """Согласовать решение всеми ролями коридора."""
        for a in proc.approvals.filter(state="pending"):
            services.resolve_decision_approval(a, self.by_role[a.role], True)

    def submit_and_approve(self, proc):
        services.submit_decision(proc, self.buyer)
        self.approve_all(proc)

    def refresh(self, *objs):
        for o in objs:
            o.refresh_from_db()


class MemoStatusTests(Base):
    def test_partial_approval_and_calculated_status(self):
        m = self.memo(self.initiator, [(self.paper, 10), (self.pens, 5)], approve=False)
        pens_item = m.items.get(line_no=2)
        services.approve_memo(m, self.head, {pens_item.pk: "нет бюджета"})
        self.refresh(m, pens_item)
        self.assertEqual(m.state, Memo.State.PARTIALLY_APPROVED)
        self.assertEqual(pens_item.status, S.REJECTED)
        self.assertEqual(m.items.get(line_no=1).status, S.APPROVED)
        self.assertEqual(m.aggregate_status[0], "new")
        self.assertIn("1 в пуле", m.summary_text())

    def test_reject_requires_reason(self):
        m = self.memo(self.initiator, [(self.paper, 10)], approve=False)
        with self.assertRaises(services.BusinessError):
            services.approve_memo(m, self.head, {m.items.first().pk: ""})

    def test_problem_status_when_overdue(self):
        m = self.memo(self.initiator, [(self.paper, 10)])
        MemoItem.objects.filter(memo=m).update(required_date=timezone.localdate() - timedelta(days=1))
        services.run_periodic()  # статус «Проблемная» зависит от даты и обновляется периодической задачей
        m = Memo.objects.get(pk=m.pk)
        self.assertEqual(m.aggregate_status[0], "problem")


class ProcurementFlowTests(Base):
    def test_full_cycle_consolidation_partial_split(self):
        m1 = self.memo(self.initiator, [(self.paper, 50), (self.pens, 10)])
        m2 = self.memo(self.initiator2, [(self.paper, 50)])
        items = list(MemoItem.objects.filter(memo__in=[m1, m2]))
        self.assertEqual(len(services.pool_items()), 3)

        proc = services.create_procurement(self.buyer, items, "Канц")
        # Позиция одновременно только в одной закупке.
        with self.assertRaises(services.BusinessError):
            services.create_procurement(self.buyer, [items[0]], "Дубль")
        self.assertEqual(services.pool_items(), [])

        # Консолидация: две строки бумаги → одна сводная 100 в запросе КП.
        rows = services.consolidated_rows(proc)
        paper_row = next(r for r in rows if r["consolidated"])
        self.assertEqual(paper_row["quantity"], D(100))

        services.add_rfqs(proc, [self.s1, self.s2], self.buyer)
        services.mark_rfqs_sent(proc, self.buyer)
        self.quote(proc, self.s1, {self.paper: 2400, self.pens: 1500}, offered={self.pens: D(6)})
        self.quote(proc, self.s2, {self.paper: 2300, self.pens: None})
        self.refresh(proc)
        self.assertEqual(proc.status, Procurement.Status.ANALYSIS)

        cmp = services.comparison(proc)
        self.assertEqual(cmp["best_total"], D(100 * 2300 + 6 * 1500))  # ручек предложено только 6

        services.decision_preset(proc, "best", self.buyer)
        pens_line = proc.lines.get(item__nomenclature=self.pens)
        self.assertEqual(pens_line.decision_lines.get().quantity, D(6))  # поставщик даёт только 6 из 10
        self.submit_and_approve(proc)
        self.refresh(proc)
        self.assertEqual(proc.decision_state, Procurement.DecisionState.APPROVED)
        pens_item = pens_line.item
        pens_item.refresh_from_db()
        self.assertEqual(pens_item.status, S.SUPPLIER_SELECTED)

        contracts = services.create_contracts(proc, self.buyer, {})
        self.assertEqual({c.supplier for c in contracts}, {self.s1, self.s2})
        c2 = next(c for c in contracts if c.supplier == self.s2)
        # Один договор — позиции двух разных СЗ.
        self.assertEqual(c2.memos.count(), 2)

        pens_item.refresh_from_db()
        self.assertEqual(pens_item.status, S.PARTIALLY_CONTRACTED)
        self.assertEqual(pens_item.remaining_qty, D(4))
        self.assertIn("остаток от закупки", pens_item.pool_note)
        self.assertEqual([i.pk for i in services.pool_items()], [pens_item.pk])

        # Исполнение: поступление по ID позиции → уведомление инициатору только по его позициям.
        for st in (Contract.Status.SIGNING, Contract.Status.SIGNED):
            services.set_contract_status(c2, st, self.buyer)
        line_m1 = c2.lines.get(item__memo=m1)
        services.register_receipt(c2, timezone.localdate(), "ПТУ-1",
                                  [(line_m1, D(50), ReceiptLine.Match.ID, True, "", "")], self.buyer)
        line_m1.item.refresh_from_db()
        self.assertEqual(line_m1.item.status, S.DELIVERED)
        self.assertTrue(self.initiator.notifications.filter(text__contains="50 из 50").exists())
        self.assertFalse(self.initiator2.notifications.filter(text__contains="поставлено").exists())
        c2.refresh_from_db()
        self.assertEqual(c2.status, Contract.Status.EXECUTION)

    def test_split_between_suppliers_and_over_quantity(self):
        m = self.memo(self.initiator, [(self.paper, 100)])
        proc = services.create_procurement(self.buyer, list(m.items.all()), "Бумага")
        services.add_rfqs(proc, [self.s1, self.s2], self.buyer)
        q1 = self.quote(proc, self.s1, {self.paper: 2400})
        q2 = self.quote(proc, self.s2, {self.paper: 2300})
        line = proc.lines.get()
        ql1, ql2 = q1.lines.get(), q2.lines.get()
        with self.assertRaises(services.BusinessError):
            services.set_decision(proc, [(line, ql1, D(60)), (line, ql2, D(60))], self.buyer)
        services.set_decision(proc, [(line, ql1, D(60)), (line, ql2, D(60))], self.buyer, allow_over=True)
        services.set_decision(proc, [(line, ql1, D(40)), (line, ql2, D(60))], self.buyer)
        self.submit_and_approve(proc)
        contracts = services.create_contracts(proc, self.buyer, {})
        self.assertEqual(len(contracts), 2)
        item = m.items.get()
        self.assertEqual(item.contracted_qty, D(100))
        self.assertEqual(item.status, S.CONTRACTED)

    def test_threshold_approval_and_analog_confirmation(self):
        m = self.memo(self.initiator, [(self.paper, 1000)])
        proc = services.create_procurement(self.buyer, list(m.items.all()), "Много бумаги")
        services.add_rfqs(proc, [self.s1], self.buyer)
        self.quote(proc, self.s1, {self.paper: 2000}, analog={self.paper: "Бумага Снегурочка"})
        services.decision_preset(proc, "best", self.buyer)
        d = DecisionLine.objects.get(procurement=proc)
        self.assertEqual(d.analog_state, DecisionLine.AnalogState.PENDING)
        with self.assertRaises(services.BusinessError):
            services.submit_decision(proc, self.buyer)
        with self.assertRaises(services.BusinessError):
            services.confirm_analog(d, self.initiator2, True)  # не инициатор
        services.confirm_analog(d, self.initiator, True)
        services.submit_decision(proc, self.buyer)  # 2 000 000 ₸ → зелёный коридор → главный инженер
        self.refresh(proc)
        self.assertEqual(proc.decision_state, Procurement.DecisionState.ON_APPROVAL)
        self.assertEqual(proc.corridor.code, "g")
        approval = proc.approvals.get()
        self.assertEqual(approval.role, roles.CHIEF)
        with self.assertRaises(services.BusinessError):
            services.resolve_decision_approval(approval, self.buyer, True)
        services.resolve_decision_approval(approval, self.chief, True)
        self.refresh(proc)
        self.assertEqual(proc.decision_state, Procurement.DecisionState.APPROVED)
        (c,) = services.create_contracts(proc, self.buyer, {})
        self.assertIn("Снегурочка", c.lines.get().description)

    def test_unselected_items_return_to_pool_and_cancel(self):
        m = self.memo(self.initiator, [(self.paper, 10), (self.pens, 5)])
        proc = services.create_procurement(self.buyer, list(m.items.all()), "Т")
        services.add_rfqs(proc, [self.s1], self.buyer)
        self.quote(proc, self.s1, {self.paper: 100, self.pens: None})
        services.decision_preset(proc, "best", self.buyer)
        self.submit_and_approve(proc)
        pens = m.items.get(nomenclature=self.pens)
        self.assertEqual(pens.status, S.APPROVED)
        self.assertIn(pens, services.pool_items())

        proc2 = services.create_procurement(self.buyer, [pens], "Ручки")
        services.cancel_procurement(proc2, self.buyer)
        pens.refresh_from_db()
        self.assertEqual(pens.status, S.APPROVED)
        self.assertIn("отменена", pens.pool_note)


class CorridorTests(Base):
    """Коридоры согласования из макета ГРЭС."""

    def priced(self, price, qty=1, method=Procurement.Method.SINGLE):
        m = self.memo(self.initiator, [(self.paper, qty)])
        proc = services.create_procurement(self.buyer, list(m.items.all()), "Т", method)
        services.quick_price(proc, self.s1, {proc.lines.get().pk: D(price)}, self.buyer)
        proc.refresh_from_db()
        return proc

    def test_corridor_by_amount_and_parallel_roles(self):
        g = self.priced(5_000_000)
        self.assertEqual((g.corridor.code, [a.role for a in g.approvals.all()]), ("g", [roles.CHIEF]))
        y = self.priced(27_500_000)
        self.assertEqual(y.corridor.code, "y")
        self.assertEqual({a.role for a in y.approvals.all()}, {roles.CHIEF, roles.CFO, roles.BUYER})
        r = self.priced(1_850_000_000)
        self.assertEqual(r.corridor.code, "r")
        self.assertEqual(r.approvals.count(), 5)
        simple = self.priced(960_000_000, method=Procurement.Method.CONTRACT)
        self.assertEqual({a.role for a in simple.approvals.all()}, {roles.PTO, roles.CFO})
        # Срок — рабочие дни: не раньше чем через 1 день и всегда в будни.
        self.assertGreater(g.approvals.get().due_at, timezone.now())
        self.assertLess(g.approvals.get().due_at.weekday(), 5)

    def test_question_does_not_reset_and_reject_needs_reason(self):
        y = self.priced(27_500_000)
        services.resolve_decision_approval(y.approvals.get(role=roles.CFO), self.cfo, True)
        q = services.ask_question(self.chief, "Почему не Темир-Хим?", procurement=y)
        self.assertEqual(y.approvals.filter(state="approved").count(), 1)  # согласование финансов не сброшено
        with self.assertRaises(services.BusinessError):
            services.answer_question(q, self.initiator2, "не я")
        services.answer_question(q, self.buyer, "нет лицензии")
        q.refresh_from_db()
        self.assertFalse(q.is_open)
        with self.assertRaises(services.BusinessError):
            services.resolve_decision_approval(y.approvals.get(role=roles.CHIEF), self.chief, False, "")
        services.resolve_decision_approval(y.approvals.get(role=roles.CHIEF), self.chief, False, "дорого")
        y.refresh_from_db()
        self.assertEqual(y.decision_state, Procurement.DecisionState.REJECTED)

    def test_bulk_green_and_escalation(self):
        g1, g2, y = self.priced(1_000_000), self.priced(2_000_000), self.priced(10_000_000)
        self.assertEqual(services.approve_all_green(self.chief), 2)
        for p in (g1, g2):
            p.refresh_from_db()
            self.assertEqual(p.decision_state, Procurement.DecisionState.APPROVED)
        y.refresh_from_db()
        self.assertEqual(y.decision_state, Procurement.DecisionState.ON_APPROVAL)  # жёлтый массово не согласуется
        r = self.priced(100_000_000)
        DecisionApproval.objects.filter(procurement=r).update(due_at=timezone.now() - timedelta(days=1))
        res = services.run_periodic()
        self.assertGreaterEqual(res["escalations"], 1)
        self.assertTrue(self.director.notifications.filter(text__startswith="Эскалация").exists())

    def test_quick_price_only_for_single_source(self):
        m = self.memo(self.initiator, [(self.paper, 1)])
        proc = services.create_procurement(self.buyer, list(m.items.all()), "Т", Procurement.Method.RFQ)
        with self.assertRaises(services.BusinessError):
            services.quick_price(proc, self.s1, {proc.lines.get().pk: D(100)}, self.buyer)

    def test_initiator_does_not_see_prices(self):
        g = self.priced(3_000_000)
        memo = Memo.objects.get(items__procurement_lines__procurement=g)
        self.assertFalse(roles.sees_prices(self.initiator))
        self.client.force_login(self.initiator)
        html = self.client.get(f"/?f=all&sel=m{memo.pk}").content.decode()
        self.assertNotIn("3 000 000", html.replace("\xa0", " "))
        self.assertIn("скрыта", html)
        self.client.force_login(self.chief)
        html = self.client.get(f"/?f=all&sel=p{g.pk}").content.decode().replace("\xa0", " ")
        self.assertIn("3 000 000", html)

    def test_payment_tranches_and_auto_close(self):
        g = self.priced(1_000, qty=10)
        self.approve_all(g)
        (c,) = services.create_contracts(g, self.buyer, {})
        services.set_contract_status(c, Contract.Status.SIGNING, self.buyer)
        services.set_contract_status(c, Contract.Status.SIGNED, self.buyer)
        services.register_payment(c, timezone.localdate(), D(3000), "ПП-1", self.cfo)
        with self.assertRaises(services.BusinessError):
            services.register_payment(c, timezone.localdate(), D(8000), "ПП-2", self.cfo)  # больше остатка
        services.register_payment(c, timezone.localdate(), D(7000), "ПП-2", self.cfo)
        services.register_receipt(c, timezone.localdate(), "ПТУ", [(c.lines.get(), D(10), ReceiptLine.Match.ID, True, "", "")])
        c.refresh_from_db()
        g.refresh_from_db()
        self.assertEqual(c.status, Contract.Status.CLOSED)
        self.assertEqual(g.status, Procurement.Status.CLOSED)


class WithdrawalTests(Base):
    def test_withdraw_from_pool_procurement_and_contract(self):
        m = self.memo(self.initiator, [(self.paper, 10), (self.pens, 5)])
        paper, pens = m.items.get(line_no=1), m.items.get(line_no=2)
        services.withdraw_item(paper, self.initiator, "не нужно")
        paper.refresh_from_db()
        self.assertEqual(paper.status, S.WITHDRAWN)

        proc = services.create_procurement(self.buyer, [pens], "Ручки")
        services.add_rfqs(proc, [self.s1], self.buyer)
        services.mark_rfqs_sent(proc, self.buyer)
        req = services.withdraw_item(pens, self.initiator, "передумали")
        self.assertIsInstance(req, WithdrawalRequest)
        pens.refresh_from_db()
        self.assertEqual(pens.status, S.IN_PROCUREMENT)
        services.resolve_withdrawal(req, self.buyer, True)
        pens.refresh_from_db()
        self.assertEqual(pens.status, S.WITHDRAWN)
        self.assertTrue(proc.rfqs.get().needs_correction)

    def test_cannot_withdraw_contracted(self):
        m = self.memo(self.initiator, [(self.paper, 10)])
        item = m.items.get()
        proc = services.create_procurement(self.buyer, [item], "Б")
        services.add_rfqs(proc, [self.s1], self.buyer)
        self.quote(proc, self.s1, {self.paper: 100})
        services.decision_preset(proc, "best", self.buyer)
        self.submit_and_approve(proc)
        services.create_contracts(proc, self.buyer, {})
        with self.assertRaises(services.BusinessError):
            services.withdraw_item(item, self.initiator, "поздно")


class OneCTests(Base):
    def _contract(self):
        m1 = self.memo(self.initiator, [(self.paper, 50)])
        m2 = self.memo(self.initiator2, [(self.paper, 30), (self.pens, 5)])
        proc = services.create_procurement(self.buyer, list(MemoItem.objects.filter(memo__in=[m1, m2])), "Т")
        services.add_rfqs(proc, [self.s1], self.buyer)
        self.quote(proc, self.s1, {self.paper: 100, self.pens: 50})
        services.decision_preset(proc, "best", self.buyer)
        self.submit_and_approve(proc)
        (c,) = services.create_contracts(proc, self.buyer, {})
        c.number = "Д-1"
        c.save()
        services.set_contract_status(c, Contract.Status.SIGNING, self.buyer)
        services.set_contract_status(c, Contract.Status.SIGNED, self.buyer)
        return c, m1, m2

    def test_export_contains_position_ids_and_collapse(self):
        c, m1, m2 = self._contract()
        data = json.loads(onec.export_json([c], collapse=False))
        ids = {l["position_id"] for l in data["contracts"][0]["lines"]}
        self.assertEqual(ids, {f"СЗ-{m1.number}/1", f"СЗ-{m2.number}/1", f"СЗ-{m2.number}/2"})
        collapsed = json.loads(onec.export_json([c], collapse=True))
        self.assertEqual(len(collapsed["contracts"][0]["lines"]), 2)
        self.assertIn("IDПозиции", onec.export_xml([c]).decode("utf-8"))

    def test_import_receipts_by_id_collapsed_and_by_name(self):
        c, m1, m2 = self._contract()
        # Строка 1 — свёрнутая (две позиции через запятую), строка 2 — без ID, только наименование.
        csv_text = (
            "Договор;Дата;НомерДокумента;IDПозиции;Номенклатура;КодНоменклатуры;Количество\n"
            f"Д-1;01.10.2026;ПТУ-1;СЗ-{m1.number}/1,СЗ-{m2.number}/1;Бумага;P1;80\n"
            "Д-1;01.10.2026;ПТУ-1;;Ручки;;5\n"
        )
        rep = onec.import_receipts(SimpleUploadedFile("r.csv", csv_text.encode("utf-8")), self.buyer)
        self.assertEqual(rep["errors"], [])
        self.assertEqual(rep["by_id"], 2)
        self.assertEqual(rep["by_nomenclature"], 1)
        self.assertEqual(m1.items.get().delivered_qty, D(50))
        self.assertEqual(m2.items.get(line_no=1).delivered_qty, D(30))
        pens = m2.items.get(line_no=2)
        self.assertEqual(pens.delivered_qty, D(0))  # ждёт ручного подтверждения
        rl = ReceiptLine.objects.get(confirmed=False)
        services.confirm_receipt_line(rl, rl.contract_line, self.buyer)
        pens.refresh_from_db()
        self.assertEqual(pens.delivered_qty, D(5))
        c.refresh_from_db()
        self.assertEqual(c.status, Contract.Status.EXECUTED)

    def test_import_payments_proportional(self):
        c, m1, m2 = self._contract()
        body = json.dumps([{"Договор": "Д-1", "Дата": "2026-10-01", "Сумма": float(c.amount / 2)}])
        rep = onec.import_payments(SimpleUploadedFile("p.json", body.encode()), self.buyer)
        self.assertEqual(rep["payments"], 1)
        item = m1.items.get()
        self.assertEqual(item.paid_amount, D("2500.00"))  # 50 × 100 / 2


class SmokeTests(TestCase):
    """Все страницы открываются у всех ролей на демо-данных."""

    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", stdout=io.StringIO())

    def test_pages_render_for_all_users(self):
        from .models import Contract as C, Procurement as P
        urls = [reverse(n) for n in ["home", "approvals", "memo_list", "memo_create", "pool", "procurement_list",
                                     "contract_list", "integration", "reports", "notifications", "rights",
                                     "dashboard", "supplier_list", "contract_create", "refs_import",
                                     "password_change"]]
        urls += [reverse("reports") + f"?r={r}" for r in ["cycle", "savings", "suppliers"]]
        urls += [m.get_absolute_url() for m in Memo.objects.all()]
        urls += [reverse("memo_print", args=[m.pk]) for m in Memo.objects.all()]
        for p in P.objects.all():
            urls += [p.get_absolute_url() + f"?tab={t}" for t, _ in __import__("core.views.procurements", fromlist=["TABS"]).TABS]
            for rfq in p.rfqs.all():
                urls += [reverse("quote_entry", args=[p.pk, rfq.pk]), reverse("rfq_print", args=[p.pk, rfq.pk])]
            urls.append(reverse("comparison_export", args=[p.pk]))
        for c in C.objects.all():
            urls += [c.get_absolute_url() + f"?tab={t}" for t in ["spec", "receipts", "payments", "history", "specs"]]
        urls += ["/?f=all&sel=" + k for k in [f"m{m.pk}" for m in Memo.objects.all()] + [f"p{p.pk}" for p in P.objects.all()]]
        urls += ["/", "/?f=late", "/?f=done", "/?f=red", "/?f=myappr", "/?f=overdue", "/?new=1", "/?repeat=3", reverse("access_matrix")]
        for username in ["serik", "beketov", "ospanova", "ahmetov", "nurlanova", "musin", "saparov", "ermekov",
                         "akhmetova", "admin"]:
            self.client.force_login(User.objects.get(username=username))
            for url in urls:
                r = self.client.get(url)
                self.assertIn(r.status_code, (200, 302, 403), f"{username} {url} → {r.status_code}")
                if r.status_code == 200 and r.get("Content-Type", "").startswith("text/html"):
                    html = r.content.decode()
                    # У каждой POST-формы должен быть CSRF-токен (иначе в браузере 403).
                    for form in re.findall(r'<form[^>]*method="post"[^>]*>(.*?)</form>', html, re.S):
                        self.assertIn("csrfmiddlewaretoken", form, f"{username} {url}: форма без csrf_token")

    def test_portal_link_works_without_login(self):
        rfq = RFQ.objects.first()
        r = self.client.get(reverse("portal_quote", args=[rfq.token]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Коммерческое предложение")


class HttpFlowTests(Base):
    """Сквозной сценарий через HTTP — как это делает пользователь в браузере."""

    def post(self, user, url, data=None, **kw):
        self.client.force_login(user)
        r = self.client.post(url, data or {}, **kw)
        self.assertIn(r.status_code, (200, 302), f"{url} → {r.status_code}")
        if r.status_code == 302:
            follow = self.client.get(r["Location"])
            msgs = list(follow.context["messages"]) if follow.context and "messages" in follow.context else []
            errors = [str(m) for m in msgs if m.level_tag == "error"]
            self.assertEqual(errors, [], url)
        return r

    def test_end_to_end(self):
        from openpyxl import load_workbook
        today = timezone.localdate()
        # 1. Инициатор создаёт СЗ, добавляет позиции, отправляет.
        self.client.force_login(self.initiator)
        r = self.client.post(reverse("memo_create"), {"department": self.dept.pk, "justification": "нужно",
                                                      "required_date": today + timedelta(days=5)})
        memo = Memo.objects.get()
        self.assertRedirects(r, memo.get_absolute_url())
        for nom, qty in [(self.paper, "40"), (self.pens, "3")]:
            self.post(self.initiator, reverse("item_add", args=[memo.pk]),
                      {"nomenclature": nom.pk, "quantity": qty, "required_date": today + timedelta(days=5)})
        self.assertEqual(memo.items.count(), 2)
        self.post(self.initiator, reverse("memo_action", args=[memo.pk, "submit"]))
        # 2. Согласующий утверждает.
        self.post(self.head, reverse("memo_action", args=[memo.pk, "approve"]), {"comment": "ок"})
        memo.refresh_from_db()
        self.assertEqual(memo.state, Memo.State.APPROVED)
        # 3. Закупщик: пул → закупка (без названия — подставится).
        self.client.force_login(self.buyer)
        self.assertEqual(self.client.get(reverse("pool")).status_code, 200)
        self.post(self.buyer, reverse("pool_create_procurement"),
                  {"items": [i.pk for i in memo.items.all()], "target": "new", "title": "", "method": "rfq",
                   "buyer": self.buyer.pk, "kp_deadline": today + timedelta(days=3)})
        proc = Procurement.objects.get()
        act = lambda a, data=None, user=self.buyer: self.post(user, reverse("procurement_action", args=[proc.pk, a]), data)
        act("add_suppliers", {"suppliers": [self.s1.pk, self.s2.pk]})
        act("send")
        rfq1, rfq2 = proc.rfqs.order_by("supplier__name")
        lines = {l.item.nomenclature_id: l for l in proc.lines.all()}
        gp, gn = lines[self.paper.pk].pk, lines[self.pens.pk].pk
        # 4. КП: вручную, из Excel-шаблона и через портал поставщика.
        self.post(self.buyer, reverse("quote_entry", args=[proc.pk, rfq1.pk]),
                  {f"price_{gp}": "2 400,50", f"lead_{gp}": "3", f"price_{gn}": "1500", f"offered_{gn}": "2"})
        wb_resp = self.client.get(reverse("quote_template", args=[proc.pk, rfq2.pk]))
        wb = load_workbook(io.BytesIO(wb_resp.content))
        ws = wb.active
        for row in ws.iter_rows(min_row=4):
            if row[0].value == gp:
                row[5].value = 2300
                row[8].value = "Бумага Снегурочка"
            elif row[0].value == gn:
                row[7].value = "да"
        buf = io.BytesIO()
        wb.save(buf)
        self.post(self.buyer, reverse("quote_entry", args=[proc.pk, rfq2.pk]),
                  {"file": SimpleUploadedFile("kp.xlsx", buf.getvalue())})
        q2 = rfq2.quote
        self.assertTrue(q2.lines.get(procurement_line_id=gp).is_analog)
        self.assertTrue(q2.lines.get(procurement_line_id=gn).not_offered)
        self.client.logout()
        r = self.client.post(reverse("portal_quote", args=[rfq1.token]),
                             {f"price_{gp}": "2350", f"price_{gn}": "1450", f"offered_{gn}": "2"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(rfq1.quote.lines.get(procurement_line_id=gp).price, D("2350"))
        # 5. Сравнение, экспорт, решение вручную (дробление бумаги между двумя поставщиками).
        self.client.force_login(self.buyer)
        for tab in ["compare", "decision", "quotes"]:
            self.assertEqual(self.client.get(proc.get_absolute_url() + f"?tab={tab}").status_code, 200)
        self.assertEqual(self.client.get(reverse("comparison_export", args=[proc.pk])).status_code, 200)
        ql1p = rfq1.quote.lines.get(procurement_line_id=gp)
        ql2p = q2.lines.get(procurement_line_id=gp)
        ql1n = rfq1.quote.lines.get(procurement_line_id=gn)
        act("decision", {f"qty_{gp}_{ql1p.pk}": "25", f"qty_{gp}_{ql2p.pk}": "15", f"qty_{gn}_{ql1n.pk}": "2"})
        self.assertEqual(proc.decision_lines.count(), 3)
        # Аналог → нужно подтверждение инициатора.
        d = proc.decision_lines.get(quote_line=ql2p)
        self.assertEqual(d.analog_state, DecisionLine.AnalogState.PENDING)
        self.post(self.initiator, reverse("analog_confirm", args=[d.pk]), {"accept": "1"})
        act("submit_decision")
        proc.refresh_from_db()
        self.assertEqual(proc.decision_state, Procurement.DecisionState.ON_APPROVAL)
        # Согласование по коридору с главного экрана: вопрос (не сбрасывает) → ответ → согласовано.
        approval = proc.approvals.get()
        self.post(self.chief, reverse("ws_ask", args=[f"p{proc.pk}"]), {"text": "Почему два поставщика?", "next": "/"})
        q = proc.questions.get()
        self.post(self.buyer, reverse("ws_answer", args=[q.pk]), {"text": "У одного не хватило объёма"})
        self.post(self.chief, reverse("ws_comment", args=[f"p{proc.pk}"]), {"text": "Ок, согласую"})
        self.post(self.chief, reverse("ws_approve", args=[approval.pk]), {"approve": "1", "next": f"/?sel=p{proc.pk}"})
        proc.refresh_from_db()
        self.assertEqual(proc.decision_state, Procurement.DecisionState.APPROVED)
        self.assertTrue(proc.history.filter(kind="ok").exists())
        self.assertTrue(proc.history.filter(kind="comment").exists())
        # 6. Договоры: два поставщика, ручки 2 из 3 → остаток 1 в пул.
        act("contracts", {f"mode_{self.s1.pk}": "new", f"mode_{self.s2.pk}": "new", f"number_{self.s1.pk}": "Д-77"})
        self.assertEqual(Contract.objects.count(), 2)
        pens = memo.items.get(nomenclature=self.pens)
        self.assertEqual(pens.status, S.PARTIALLY_CONTRACTED)
        self.assertIn(pens, services.pool_items())
        c = Contract.objects.get(number="Д-77")
        for st in ["signing", "signed"]:
            self.post(self.buyer, reverse("contract_action", args=[c.pk, "status"]), {"status": st})
        # 7. Исполнение: поступление вручную и из 1С, оплата, выгрузка.
        line = c.lines.get(item__nomenclature=self.paper)
        self.post(self.buyer, reverse("contract_action", args=[c.pk, "receipt"]),
                  {"date": today.isoformat(), "doc": "ПТУ-9", f"qty_{line.pk}": "5"})
        self.assertEqual(line.delivered_qty, D(5))
        csv_text = f"Договор;Дата;НомерДокумента;IDПозиции;Номенклатура;КодНоменклатуры;Количество\nД-77;{today:%d.%m.%Y};ПТУ-10;{line.item.code};;;20\n"
        self.post(self.buyer, reverse("integration"), {"kind": "receipts", "file": SimpleUploadedFile("r.csv", csv_text.encode())})
        self.assertEqual(line.delivered_qty, D(25))
        self.post(self.buyer, reverse("contract_action", args=[c.pk, "payment"]), {"date": today.isoformat(), "amount": "10000", "doc": "ПП"})
        self.assertEqual(c.paid_total, D(10000))
        for fmt in ["xml", "json", "xlsx"]:
            r = self.client.post(reverse("integration"), {"kind": "export", "contracts": [c.pk], "format": fmt})
            self.assertEqual(r.status_code, 200)
        # 8. Инициатор видит уведомление только по своим позициям и может отозвать остаток ручек.
        self.assertTrue(self.initiator.notifications.filter(text__contains="поставлено").exists())
        self.post(self.initiator, reverse("item_action", args=[pens.pk, "withdraw"]), {"reason": "хватит"})
        pens.refresh_from_db()
        self.assertTrue(pens.remainder_closed)
        self.assertEqual(pens.status, S.CONTRACTED)
        # 9. Отчёты строятся (у инициатора доступа к ним нет).
        self.assertEqual(self.client.get(reverse("reports")).status_code, 403)
        self.client.force_login(self.buyer)
        for r_ in ["coverage", "cycle", "savings", "suppliers"]:
            self.assertEqual(self.client.get(reverse("reports") + f"?r={r_}&export=1").status_code, 200)


class ProductionTests(Base):
    def test_refs_import_template_roundtrip(self):
        from openpyxl import load_workbook
        from .refs_import import build_template, import_workbook
        wb = load_workbook(io.BytesIO(build_template()))
        self.assertEqual(import_workbook(build_template()).created, {})  # пустой шаблон ничего не создаёт
        def put(sheet, *rows):
            ws = wb[sheet]
            for r in rows:
                ws.append(list(r))
        put("Подразделения", ("Финансы", "fin_head"))
        put("Пользователи", ("fin_head", "Ахметова", "Алия", "a@x.kz", "Финансы", "Руководитель", "ГРЭС-9", "Руководитель подразделения"),
            ("zakup1", "Касымов", "Данияр", "", "", "", "", "Закупки (ОМТС), Инициатор"))
        put("Категории", ("Канцтовары",))
        put("Номенклатура", ("Бумага А4", "пач", "Канцтовары", "00-1"))
        put("Поставщики", ("ТОО Альфа", "123456789012", "", "", "", "Канцтовары", ""))
        put("Рамочные договоры", ("РД-1", "01.01.2026", "123456789012", "31.12.2026"))
        put("Станции", ("ГРЭС-9",))
        put("Коридоры согласования", ("y", "Жёлтый", "40 000 000", "Главный инженер, Финансы (ФЭО)", "3", ""))
        put("Статьи бюджета")
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()

        dry = import_workbook(data, dry_run=True)
        self.assertEqual(dry.errors, [])
        self.assertFalse(User.objects.filter(username="zakup1").exists())  # проверка ничего не пишет

        rep = import_workbook(data)
        self.assertEqual(rep.errors, [])
        self.assertEqual(Department.objects.get(name="Финансы").head.username, "fin_head")
        u = User.objects.get(username="zakup1")
        self.assertTrue(roles.has_role(u, roles.BUYER))
        self.assertEqual(len(rep.passwords), 2)
        self.assertTrue(Contract.objects.filter(number="РД-1", kind=Contract.Kind.FRAMEWORK).exists())
        self.assertEqual(Corridor.objects.get(code="y").max_amount, D(40000000))
        self.assertTrue(Station.objects.filter(name="ГРЭС-9").exists())
        # Повторная загрузка — обновление, без дублей и без новых паролей.
        rep2 = import_workbook(data)
        self.assertEqual(rep2.created.get("Пользователи"), None)
        self.assertEqual(rep2.passwords, [])
        self.assertEqual(Supplier.objects.filter(bin="123456789012").count(), 1)

    def test_refs_import_errors_rollback(self):
        from openpyxl import load_workbook
        from .refs_import import build_template, import_workbook
        wb = load_workbook(io.BytesIO(build_template()))
        wb["Пользователи"].append(["ok_user", "А", "Б", "", "", "", "", "Инициатор"])
        wb["Пользователи"].append(["bad user", "А", "Б", "", "", "", "", "Инициатор"])
        wb["Пользователи"].append(["u3", "А", "Б", "", "Нет такого", "", "", "Космонавт"])
        buf = io.BytesIO()
        wb.save(buf)
        rep = import_workbook(buf.getvalue())
        self.assertEqual(len(rep.errors), 2)
        self.assertIn("строка 3", rep.errors[0])
        self.assertFalse(User.objects.filter(username="ok_user").exists())  # при ошибках ничего не записано

    def test_onec_reimport_is_idempotent_and_folder_exchange(self):
        import tempfile
        from pathlib import Path
        m = self.memo(self.initiator, [(self.paper, 10)])
        proc = services.create_procurement(self.buyer, list(m.items.all()), "Т")
        services.add_rfqs(proc, [self.s1], self.buyer)
        self.quote(proc, self.s1, {self.paper: 100})
        services.decision_preset(proc, "best", self.buyer)
        self.submit_and_approve(proc)
        (c,) = services.create_contracts(proc, self.buyer, {})
        c.number = "Д-5"
        c.save()
        services.set_contract_status(c, Contract.Status.SIGNING, self.buyer)
        services.set_contract_status(c, Contract.Status.SIGNED, self.buyer)
        csv_text = f"Договор;Дата;НомерДокумента;IDПозиции;Номенклатура;КодНоменклатуры;Количество\nД-5;01.10.2026;ПТУ-7;{m.items.get().code};;;4\n"
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            call_command("onec_exchange", dir=tmp, stdout=io.StringIO())
            self.assertEqual(len(list((base / "out").glob("contracts_*.xml"))), 1)
            (base / "in" / "receipts_1.csv").write_text(csv_text, encoding="utf-8")
            (base / "in" / "receipts_2.csv").write_text(csv_text, encoding="utf-8")  # 1С прислала повторно
            (base / "in" / "payments_bad.csv").write_text("Договор;Сумма\nНЕТ;1\n", encoding="utf-8")
            call_command("onec_exchange", dir=tmp, stdout=io.StringIO(), stderr=io.StringIO())
            self.assertEqual(len(list((base / "archive").iterdir())), 2)
            self.assertEqual(len(list((base / "error").glob("*.log"))), 1)
            # Второй запуск не выгружает договор повторно.
            call_command("onec_exchange", dir=tmp, stdout=io.StringIO())
            self.assertEqual(len(list((base / "out").glob("contracts_*.xml"))), 1)
        item = m.items.get()
        self.assertEqual(item.delivered_qty, D(4))  # не 8

    def test_login_lockout_and_health(self):
        from django.core.cache import cache
        call_command("createcachetable", stdout=io.StringIO())
        cache.clear()
        for _ in range(5):
            self.client.post(reverse("login"), {"username": "buyer", "password": "wrong"})
        r = self.client.post(reverse("login"), {"username": "buyer", "password": "x"})  # верный пароль, но блок
        self.assertNotIn("_auth_user_id", self.client.session)
        self.assertContains(r, "Слишком много")
        cache.clear()
        r = self.client.post(reverse("login"), {"username": "buyer", "password": "x"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.client.get("/health/").json()["status"], "ok")

    def test_concurrent_numbering_uses_counter(self):
        m1 = self.memo(self.initiator, [(self.paper, 1)], approve=False)
        Memo.objects.filter(pk=m1.pk).delete()  # «дыра» в номерах не приводит к повтору
        m2 = self.memo(self.initiator, [(self.paper, 1)], approve=False)
        self.assertGreater(m2.number, m1.number)


class WorkspaceFlowTests(Base):
    """Сценарий макета ГРЭС через главный экран."""

    def setUp(self):
        super().setUp()
        st = Station.objects.create(name="ГРЭС-1")
        Profile = __import__("core.models", fromlist=["Profile"]).Profile
        Profile.objects.create(user=self.initiator, department=self.dept, station=st)

    def test_mockup_scenario(self):
        from django.core.files.uploadedfile import SimpleUploadedFile as F
        c = self.client
        # 1. Инициатор: новая потребность из окна (позиции без цен) + скан СЗ, сразу «Отправить».
        c.force_login(self.initiator)
        r = c.post(reverse("ws_new_memo"), {
            "name": "Комплектующие ПЭН-2", "item_name": ["Торцевое уплотнение", "Бумага А4"], "item_qty": ["2", "10"],
            "item_unit": ["компл.", "пач"], "item_spec": ["по чертежу", ""], "why": "течь по валу",
            "deliv": (timezone.localdate() + timedelta(days=20)).isoformat(), "send": "1",
            "file": F("СЗ.txt", "скан".encode()),
        })
        memo = Memo.objects.get(title="Комплектующие ПЭН-2")
        self.assertRedirects(r, f"/?sel=m{memo.pk}&f=all", fetch_redirect_response=False)
        self.assertEqual(memo.state, Memo.State.ON_APPROVAL)
        self.assertEqual(memo.items.get(line_no=2).nomenclature, self.paper)  # позиция связалась со справочником
        self.assertEqual(memo.attachments.count(), 1)
        self.assertEqual(c.get(f"/?sel=m{memo.pk}").status_code, 200)
        # 2. Руководитель: «Мои задачи» → согласовать.
        c.force_login(self.head)
        self.assertContains(c.get("/"), "Согласовать СЗ")
        c.post(reverse("memo_action", args=[memo.pk, "approve"]), {"next": f"/?sel=m{memo.pk}"})
        memo.refresh_from_db()
        self.assertTrue(memo.is_approved)
        # 3. Закупщик: из пула в закупку «из одного источника» → быстрая расценка.
        proc = services.create_procurement(self.buyer, list(memo.items.all()), "ПЭН-2", Procurement.Method.SINGLE)
        c.force_login(self.buyer)
        data = {"supplier": self.s1.pk, "method": "single", "pay": "100% после поставки", "next": f"/?sel=p{proc.pk}",
                "files": [F("КП.txt", "цены".encode())]}
        for line in proc.lines.all():
            data[f"price_{line.pk}"] = "150 000" if line.item.line_no == 1 else "2500"
        c.post(reverse("ws_quick_price", args=[proc.pk]), data)
        proc.refresh_from_db()
        self.assertEqual(proc.corridor.code, "g")
        self.assertEqual(proc.decision_total, D(325000))
        # Инициатор не видит сумму и не может открыть КП.
        quote_file = proc.attachments.get()
        c.force_login(self.initiator)
        self.assertEqual(c.get(reverse("attachment", args=[quote_file.pk])).status_code, 403)
        # 4. Главный инженер: «Согласовать все зелёные».
        c.force_login(self.chief)
        c.post(reverse("ws_approve_green"))
        proc.refresh_from_db()
        self.assertEqual(proc.decision_state, Procurement.DecisionState.APPROVED)
        # 5. Закупщик: оформить договор, подписать, прикрепить скан.
        c.force_login(self.buyer)
        c.post(reverse("procurement_action", args=[proc.pk, "contracts"]), {f"mode_{self.s1.pk}": "new", "next": "/"})
        contract = Contract.objects.get(procurement=proc)
        for st in ("signing", "signed"):
            c.post(reverse("contract_action", args=[contract.pk, "status"]), {"status": st, "next": "/"})
        c.post(reverse("ws_attach", args=[f"c{contract.pk}"]), {"kind": "contract", "file": F("Договор.txt", b"x")})
        self.assertEqual(contract.attachments.count(), 1)
        # 6. Финансы: транш 30%, затем остаток.
        c.force_login(self.cfo)
        self.assertContains(c.get("/"), "Оплата")
        c.post(reverse("contract_action", args=[contract.pk, "payment"]),
               {"date": timezone.localdate().isoformat(), "amount": "97500", "doc": "ПП-1", "next": "/"})
        c.post(reverse("contract_action", args=[contract.pk, "payment"]),
               {"date": timezone.localdate().isoformat(), "amount": "227500", "doc": "ПП-2", "next": "/"})
        self.assertEqual(contract.paid_total, D(325000))
        # Выгрузка реестра в Excel.
        self.assertEqual(c.get(reverse("ws_export")).status_code, 200)

    def test_demo_switch_only_in_demo_mode(self):
        from django.test import override_settings
        with override_settings(DEMO_MODE=False):
            self.assertEqual(self.client.post(reverse("demo_switch"), {"username": "chief"}).status_code, 403)
        with override_settings(DEMO_MODE=True):
            self.client.post(reverse("demo_switch"), {"username": "chief"})
            self.assertEqual(int(self.client.session["_auth_user_id"]), self.chief.pk)
